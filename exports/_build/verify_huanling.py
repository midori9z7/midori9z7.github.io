# -*- coding: utf-8 -*-
"""校验合集：正文 token 序列是否与 5 个源文件按顺序拼接一致。"""
import html as htmllib
import os
import re

from docx import Document

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DOCX = os.path.join(ROOT, "exports", "诗歌", "幻灵（合集）.docx")
SRC = [
    "posts/poems/2026/幻灵，第4号.html",
    "posts/poems/2026/幻灵，第6号.html",
    "posts/poems/2026/幻灵，第8号.html",
    "posts/poems/2026/幻灵，第13号.html",
    "posts/poems/2026/幻灵，第14号.html",
]
TOKEN = re.compile(r"[\u3400-\u9fff]|[A-Za-z0-9]+")


def tokens(s):
    return TOKEN.findall(s)


exp = []
for rel in SRC:
    text = open(os.path.join(ROOT, rel), encoding="utf-8").read()
    title = re.search(r'<h1 class="title">(.*?)</h1>', text, re.S).group(1)
    verse = re.search(r'<pre class="verse">(.*?)</pre>', text, re.S).group(1)
    exp += tokens(htmllib.unescape(title)) + tokens(htmllib.unescape(verse))

d = Document(DOCX)
paras = [p.text for p in d.paragraphs]
got = tokens("\n".join(paras[1:]))  # 去掉合集大标题「幻灵」
print("paragraphs:", len(paras), "| tokens got:", len(got), "exp:", len(exp))
if got == exp:
    print("content: OK")
else:
    i = 0
    while i < min(len(got), len(exp)) and got[i] == exp[i]:
        i += 1
    print("content: DIFF at", i, "got", "".join(got[i:i + 20]), "exp", "".join(exp[i:i + 20]))
print("first lines:", paras[:3])
breaks = [p.text for p in d.paragraphs if p.paragraph_format.page_break_before]
print("page-break paragraphs:", breaks)
